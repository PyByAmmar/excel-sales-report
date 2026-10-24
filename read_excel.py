from openpyxl import load_workbook, Workbook

def read_sales():
    totals = {}
    skipped = []
    processed = 0
    wb = load_workbook("sales_report.xlsx")
    sheet = wb.active
    
    for row_number, row in enumerate( sheet.iter_rows(min_row=2, values_only=True), start=2):
        
        if row[1] is  None and row[2] is None:
            continue

        try:

            category = row[1].strip().title()
            amount = float( row[2])
            
            
        except ValueError:
            skipped.append(f"row {row_number} (invalid amount)")
            continue
        except TypeError:
            skipped.append(f"row {row_number} (empty amount)")
            continue
        except AttributeError:
            skipped.append(f"row {row_number} (missing category)")
            continue
        except PermissionError:
            continue
            
        totals[category] = totals.get(category, 0) +amount
        processed += 1
       
    return totals, skipped, processed

def write_report(totals):
    wb = Workbook()
    sheet = wb.active
    sheet.append(["category", "total"])
    for category, total in totals.items():
        sheet.append([category, total])
        
    wb.save("report.xlsx")
    
try:
    totals, skipped, processed = read_sales()
    write_report(totals)
    print("Done")
    print("Processed", processed, "rows")
    
    if skipped:
        print("Skipped", len(skipped), "rows:", ", ".join(skipped))

except FileNotFoundError:
    print("File not found.")

